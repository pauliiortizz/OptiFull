"""Tests de frontend/api/reportes_formatos.py y reportes_graficos.py: generacion de los archivos de
Reportes 2.0 (CSV, PNG, Excel y PDF) a partir de metricas armadas a mano, sin base de datos."""
import csv
import io
import zipfile
from datetime import date

import pytest
from openpyxl import load_workbook

from api import reportes_formatos as rf
from api.reportes_datos import Filtros
from api.reportes_graficos import grafico_png

PNG = b"\x89PNG"
FILTROS = Filtros(desde=date(2026, 10, 1), hasta=date(2026, 10, 7), camaras=[2])
BARRAS = {"id": "g1", "tipo": "bar", "titulo": "Personas por día", "labels": ["Lun", "Mar"], "valores": [3, 5],
          "ylabel": "personas"}


def metrica(id_="m1", tablas=(), graficos=(), resumen=(("Promedio", 4),), notas=()):
    return {"id": id_, "titulo": f"Métrica {id_}", "categoria": "Personas / Afluencia", "resumen": list(resumen),
            "tablas": list(tablas), "graficos": list(graficos), "notas": list(notas)}


def tabla(nombre="Flujo por día", filas=(("Lun", 3), ("Mar", None)), solo_datos=False):
    return {"nombre": nombre, "columnas": ["Día", "Personas"], "filas": [list(f) for f in filas],
            "solo_datos": solo_datos}


# ── Utilidades ────────────────────────────────────────────────────────────────

@pytest.mark.parametrize("texto, esperado", [
    ("Permanencia por día", "permanencia_por_dia"),
    ("  Góndolas / Caja  ", "gondolas_caja"),
    ("Ñandú 2026!", "nandu_2026"),
])
def test_slug_para_nombres_de_archivo(texto, esperado):
    assert rf._slug(texto) == esperado


def test_texto_para_pdf_reemplaza_caracteres_fuera_de_latin1():
    assert rf._t("A → B — ≤ 5 … “ok”") == 'A -> B - <= 5 ... "ok"'
    assert rf._t("ñandú 😀") == "ñandú ?"   # latin-1 conserva acentos; lo demas se reemplaza


# ── CSV ───────────────────────────────────────────────────────────────────────

def test_csv_sin_tablas_devuelve_none():
    assert rf.generar_csv([metrica()], FILTROS, "Todas") is None


def test_csv_una_tabla_devuelve_csv_directo_con_bom_para_excel():
    contenido, mime, ext = rf.generar_csv([metrica(tablas=[tabla()])], FILTROS, "Todas")
    assert (mime, ext) == ("text/csv", "csv")
    assert contenido.startswith(b"\xef\xbb\xbf")   # utf-8-sig: Excel abre bien los acentos
    filas = list(csv.reader(io.StringIO(contenido.decode("utf-8-sig"))))
    assert filas == [["Día", "Personas"], ["Lun", "3"], ["Mar", ""]]


def test_csv_varias_tablas_devuelve_zip_con_parametros():
    metricas = [metrica("a", tablas=[tabla("Uno")]), metrica("b", tablas=[tabla("Dos")])]
    contenido, mime, ext = rf.generar_csv(metricas, FILTROS, "Cámara 2")
    assert (mime, ext) == ("application/zip", "zip")
    with zipfile.ZipFile(io.BytesIO(contenido)) as z:
        assert sorted(z.namelist()) == ["a__uno.csv", "b__dos.csv", "parametros.txt"]
        parametros = z.read("parametros.txt").decode()
    assert "Período: 01/10/2026 al 07/10/2026" in parametros
    assert "Cámaras: Cámara 2" in parametros


# ── PNG y graficos ────────────────────────────────────────────────────────────

def test_png_sin_graficos_devuelve_none():
    assert rf.generar_png([metrica()], FILTROS, "Todas") is None


def test_png_un_grafico_devuelve_la_imagen():
    contenido, mime, ext = rf.generar_png([metrica(graficos=[BARRAS])], FILTROS, "Todas")
    assert (mime, ext) == ("image/png", "png")
    assert contenido.startswith(PNG)


def test_png_varios_graficos_devuelve_zip():
    contenido, _, ext = rf.generar_png([metrica(graficos=[BARRAS, {**BARRAS, "id": "g2"}])], FILTROS, "Todas")
    assert ext == "zip"
    with zipfile.ZipFile(io.BytesIO(contenido)) as z:
        assert sorted(z.namelist()) == ["m1__g1.png", "m1__g2.png", "parametros.txt"]


@pytest.mark.parametrize("spec", [
    BARRAS,
    {**BARRAS, "horizontal": True},
    {**BARRAS, "labels": [str(i) for i in range(20)], "valores": list(range(20))},   # >8 rota etiquetas, >14 sin rotulos
    {**BARRAS, "horizontal": True, "ylabel": None},
    {"id": "l", "tipo": "line", "titulo": "Línea", "labels": [f"d{i}" for i in range(25)], "valores": list(range(25)),
     "ylabel": "personas"},
    {"id": "h", "tipo": "heat", "titulo": "Matriz", "filas": ["Lun", "Mar"],
     "matriz": [[0] * 23 + [5], [2] * 24], "con_datos": [[False] * 23 + [True], [True] * 24],
     "picos": [(0, 23, 23)]},
], ids=["barras", "horizontal", "muchas-etiquetas", "horizontal-sin-eje", "linea", "matriz"])
def test_cada_tipo_de_grafico_genera_un_png(spec):
    assert grafico_png(spec).startswith(PNG)


def test_tipo_de_grafico_desconocido_falla():
    with pytest.raises(KeyError):
        grafico_png({**BARRAS, "tipo": "torta"})


# ── Excel ─────────────────────────────────────────────────────────────────────

def test_xlsx_tiene_resumen_una_hoja_por_tabla_y_hoja_de_graficos():
    metricas = [metrica(tablas=[tabla("Flujo: por día?")], graficos=[BARRAS], notas=["Nota de prueba"])]
    contenido, _, ext = rf.generar_xlsx(metricas, FILTROS, "Todas")
    assert ext == "xlsx"
    wb = load_workbook(io.BytesIO(contenido))
    assert wb.sheetnames == ["Resumen", "Flujo por día", "Gráficos e imágenes"]   # sin caracteres invalidos
    resumen = [c.value for c in wb["Resumen"]["A"] if c.value]
    assert "Nota de prueba" in resumen
    assert [c.value for c in wb["Flujo por día"][1]] == ["Día", "Personas"]


def test_xlsx_nombres_de_hoja_repetidos_y_largos():
    largo = "Permanencia promedio por zona del local"   # > 31 caracteres (limite de Excel)
    metricas = [metrica(tablas=[tabla(largo), tabla(largo), tabla(largo)])]
    wb = load_workbook(io.BytesIO(rf.generar_xlsx(metricas, FILTROS, "Todas")[0]))
    hojas = wb.sheetnames[1:]
    assert hojas == [largo[:31], largo[:27] + " (2)", largo[:27] + " (3)"]
    assert all(len(h) <= 31 for h in hojas)


# ── PDF ───────────────────────────────────────────────────────────────────────

def test_pdf_completo_con_resumen_tablas_graficos_y_notas():
    muchas = [(f"día {i}", i) for i in range(rf.PDF_MAX_FILAS + 5)]   # supera el limite de filas del PDF
    metricas = [
        metrica("a", tablas=[tabla(filas=muchas), tabla("Solo datos", solo_datos=True)], graficos=[BARRAS],
                notas=["Nota → con flecha"]),
        metrica("b", resumen=()),   # sin datos: "Sin datos en el período seleccionado."
    ]
    contenido, mime, ext = rf.generar_pdf(metricas, FILTROS, "Todas")
    assert (mime, ext) == ("application/pdf", "pdf")
    assert contenido.startswith(b"%PDF")
    assert len(contenido) > 5_000


def test_generadores_disponibles():
    assert set(rf.GENERADORES) == {"pdf", "xlsx", "csv", "png"}
