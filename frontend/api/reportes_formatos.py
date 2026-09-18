"""Generadores de archivo de Reportes 2.0: PDF, Excel, CSV (o ZIP de CSV) y
PNG (o ZIP de PNG). Todos reciben la misma lista de resultados de
reportes_datos.recolectar() y devuelven (bytes, mimetype, extension).
"""
import csv
import io
import re
import zipfile
from datetime import datetime

from .reportes_graficos import BRAND_DEEP, grafico_png

PDF_MAX_FILAS = 25   # filas por tabla en el PDF (el detalle completo va en Excel/CSV)
MIME = {
    'pdf':  'application/pdf',
    'xlsx': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    'csv':  'text/csv',
    'png':  'image/png',
    'zip':  'application/zip',
}


def _slug(texto):
    import unicodedata
    s = unicodedata.normalize('NFKD', texto).encode('ascii', 'ignore').decode()
    return re.sub(r'[^a-z0-9]+', '_', s.lower()).strip('_')


def _zip(archivos):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, 'w', zipfile.ZIP_DEFLATED) as z:
        for nombre, contenido in archivos:
            z.writestr(nombre, contenido)
    return buf.getvalue()


def _parametros_txt(f, camaras_txt, metricas):
    return '\n'.join([
        'OptiFull - Reporte',
        f'Generado: {datetime.now():%d/%m/%Y %H:%M}',
        f'Período: {f.texto_periodo()}',
        f'Cámaras: {camaras_txt}',
        'Métricas: ' + ', '.join(m['titulo'] for m in metricas),
    ]) + '\n'


def _jpeg(png):
    """Recomprime a JPEG para PDF/Excel: los mapas de calor llevan la foto del
    local de fondo y como PNG pesan varios MB cada uno."""
    from PIL import Image
    buf = io.BytesIO()
    Image.open(io.BytesIO(png)).convert('RGB').save(buf, 'JPEG', quality=85, optimize=True)
    return buf.getvalue()


def _imagenes(metricas):
    """(nombre_archivo, png) de todos los graficos e imagenes de las metricas."""
    for m in metricas:
        for g in m['graficos']:
            yield f"{m['id']}__{g['id']}.png", grafico_png(g)
        for img in m['imagenes']:
            yield f"{m['id']}__{img['nombre']}.png", img['png']


# ── CSV / PNG ──────────────────────────────────────────────────────────────

def generar_csv(metricas, f, camaras_txt):
    """Una sola tabla -> CSV directo; varias -> ZIP con un CSV por tabla (mas
    comodo para analisis que mezclar tablas de forma distinta en un archivo)."""
    archivos = []
    for m in metricas:
        for t in m['tablas']:
            buf = io.StringIO()
            w = csv.writer(buf)
            w.writerow(t['columnas'])
            w.writerows(t['filas'])
            # utf-8-sig para que Excel abra bien los acentos
            archivos.append((f"{m['id']}__{_slug(t['nombre'])}.csv", buf.getvalue().encode('utf-8-sig')))
    if not archivos:
        return None
    if len(archivos) == 1:
        return archivos[0][1], MIME['csv'], 'csv'
    archivos.append(('parametros.txt', _parametros_txt(f, camaras_txt, metricas).encode('utf-8')))
    return _zip(archivos), MIME['zip'], 'zip'


def generar_png(metricas, f, camaras_txt):
    imgs = list(_imagenes(metricas))
    if not imgs:
        return None
    if len(imgs) == 1:
        return imgs[0][1], MIME['png'], 'png'
    imgs.append(('parametros.txt', _parametros_txt(f, camaras_txt, metricas).encode('utf-8')))
    return _zip(imgs), MIME['zip'], 'zip'


# ── Excel ──────────────────────────────────────────────────────────────────

def generar_xlsx(metricas, f, camaras_txt):
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XlImage
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from PIL import Image

    relleno = PatternFill('solid', fgColor=BRAND_DEEP.lstrip('#'))
    cabecera = Font(bold=True, color='FFFFFF')
    wb = Workbook()

    ws = wb.active
    ws.title = 'Resumen'
    ws['A1'] = 'OptiFull - Reporte'
    ws['A1'].font = Font(bold=True, size=16, color=BRAND_DEEP.lstrip('#'))
    for i, (k, v) in enumerate([('Generado', f'{datetime.now():%d/%m/%Y %H:%M}'),
                                ('Período', f.texto_periodo()), ('Cámaras', camaras_txt)], start=2):
        ws.cell(i, 1, k).font = Font(bold=True)
        ws.cell(i, 2, v)
    fila = 6
    for m in metricas:
        ws.cell(fila, 1, m['titulo']).font = Font(bold=True, size=12, color=BRAND_DEEP.lstrip('#'))
        fila += 1
        for k, v in m['resumen']:
            ws.cell(fila, 1, k)
            ws.cell(fila, 2, v).alignment = Alignment(horizontal='left')
            fila += 1
        for nota in m['notas']:
            ws.cell(fila, 1, nota).font = Font(italic=True, color='64748B')
            fila += 1
        fila += 1
    ws.column_dimensions['A'].width = 44
    ws.column_dimensions['B'].width = 40

    usados = {'Resumen'}
    for m in metricas:
        for t in m['tablas']:
            base = re.sub(r'[\[\]:*?/\\]', '', t['nombre'])[:31]
            nombre, n = base, 2
            while nombre in usados:
                sufijo = f' ({n})'
                nombre, n = base[:31 - len(sufijo)] + sufijo, n + 1
            usados.add(nombre)
            hoja = wb.create_sheet(nombre)
            hoja.append(t['columnas'])
            for celda in hoja[1]:
                celda.fill, celda.font = relleno, cabecera
                celda.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            for r in t['filas']:
                hoja.append(r)
            hoja.freeze_panes = 'A2'
            for i, col in enumerate(t['columnas'], start=1):
                ancho = max([len(str(col))] + [len(str(r[i - 1])) for r in t['filas'] if r[i - 1] is not None])
                hoja.column_dimensions[get_column_letter(i)].width = min(max(ancho + 2, 10), 50)

    imgs = list(_imagenes(metricas))
    if imgs:
        hoja = wb.create_sheet('Gráficos e imágenes')
        y = 1
        for nombre, png in imgs:
            png = _jpeg(png)
            hoja.cell(y, 1, nombre.removesuffix('.png')).font = Font(bold=True, color='64748B')
            w0, h0 = Image.open(io.BytesIO(png)).size
            xl = XlImage(io.BytesIO(png))
            xl.width, xl.height = 760, int(760 * h0 / w0)
            hoja.add_image(xl, f'A{y + 1}')
            y += int(xl.height / 20) + 4   # ~20 px por fila por defecto

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), MIME['xlsx'], 'xlsx'


# ── PDF ────────────────────────────────────────────────────────────────────

_REEMPLAZOS = {'—': '-', '–': '-', '≤': '<=', '≥': '>=', '→': '->', '•': '-', '…': '...', '’': "'", '“': '"', '”': '"'}


def _t(s):
    """fpdf con fuentes core (Helvetica) solo soporta latin-1."""
    s = str(s)
    for k, v in _REEMPLAZOS.items():
        s = s.replace(k, v)
    return s.encode('latin-1', 'replace').decode('latin-1')


def generar_pdf(metricas, f, camaras_txt):
    from fpdf import FPDF
    from fpdf.enums import Align, XPos, YPos
    from fpdf.fonts import FontFace

    marca = tuple(int(BRAND_DEEP[i:i + 2], 16) for i in (1, 3, 5))
    ahora = datetime.now()

    class PDF(FPDF):
        def header(self):
            self.set_font('Helvetica', 'B', 9)
            self.set_text_color(*marca)
            self.cell(0, 6, 'OptiFull', align='L')
            self.set_font('Helvetica', '', 8)
            self.set_text_color(120, 120, 120)
            self.cell(0, 6, ahora.strftime('%d/%m/%Y  %H:%M'), align='R', new_x=XPos.LMARGIN, new_y=YPos.NEXT)
            self.set_draw_color(225, 225, 235)
            self.line(self.l_margin, self.get_y() + 1, self.w - self.r_margin, self.get_y() + 1)
            self.ln(5)

        def footer(self):
            self.set_y(-12)
            self.set_font('Helvetica', 'I', 8)
            self.set_text_color(150, 150, 150)
            self.cell(0, 8, f'Página {self.page_no()}/{{nb}}', align='C')

    pdf = PDF(format='A4')
    pdf.set_margins(14, 14, 14)
    pdf.set_auto_page_break(True, margin=16)
    pdf.alias_nb_pages()
    pdf.add_page()
    ancho = pdf.w - pdf.l_margin - pdf.r_margin

    def nl(alto=6):
        pdf.ln(alto)

    def estado_tabla():
        """fpdf2 hereda fuente/colores del ultimo texto dibujado: se resetean
        antes de cada tabla para que no salgan azules/negrita."""
        pdf.set_font('Helvetica', '', 8)
        pdf.set_text_color(40, 40, 50)
        pdf.set_fill_color(255, 255, 255)
        pdf.set_draw_color(215, 215, 228)

    def anchos(columnas, filas):
        """Ancho de columna proporcional al contenido (acotado), para que
        'Producto' no se parta en 4 renglones al lado de un 'SKU' corto."""
        largos = [min(max([len(str(c))] + [len(str(r[i])) for r in filas[:PDF_MAX_FILAS] if r[i] is not None]), 34)
                  for i, c in enumerate(columnas)]
        return tuple(max(l, 9) for l in largos)

    def texto(s, tam=10, estilo='', color=(40, 40, 50), alto=6):
        pdf.set_font('Helvetica', estilo, tam)
        pdf.set_text_color(*color)
        pdf.multi_cell(0, alto, _t(s), new_x=XPos.LMARGIN, new_y=YPos.NEXT)

    # Portada + resumen ejecutivo
    pdf.set_font('Helvetica', 'B', 22)
    pdf.set_text_color(*marca)
    pdf.cell(0, 12, 'Reporte de métricas', new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    texto(f'Período: {f.texto_periodo()}', 10, color=(90, 90, 100))
    texto(f'Cámaras: {camaras_txt}', 10, color=(90, 90, 100))
    nl(4)
    texto('Resumen ejecutivo', 13, 'B', (30, 30, 40), 8)
    for m in metricas:
        texto(m['titulo'], 10.5, 'B', marca)
        if m['resumen']:
            estado_tabla()
            with pdf.table(col_widths=(ancho * 0.6, ancho * 0.4), first_row_as_headings=False,
                           borders_layout='HORIZONTAL_LINES', line_height=5.5, text_align=('LEFT', 'RIGHT'),
                           padding=1) as tabla:
                for k, v in m['resumen']:
                    fila = tabla.row()
                    fila.cell(_t(k))
                    fila.cell(_t(v))
        else:
            texto('Sin datos en el período seleccionado.', 9, 'I', (130, 130, 140))
        nl(3)

    # Detalle por metrica
    for m in metricas:
        pdf.add_page()
        pdf.set_fill_color(*marca)
        pdf.set_text_color(255, 255, 255)
        pdf.set_font('Helvetica', 'B', 12)
        pdf.cell(0, 9, '  ' + _t(f"{m['categoria']}  |  {m['titulo']}"), fill=True,
                 new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        nl(4)
        for k, v in m['resumen']:
            pdf.set_font('Helvetica', '', 10)
            pdf.set_text_color(90, 90, 100)
            pdf.cell(ancho * 0.55, 6, _t(k))
            pdf.set_font('Helvetica', 'B', 10)
            pdf.set_text_color(30, 30, 40)
            pdf.cell(0, 6, _t(v), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        if m['resumen']:
            nl(3)
        for g in m['graficos']:
            pdf.image(io.BytesIO(_jpeg(grafico_png(g))), w=ancho * 0.92, x=Align.C)
            nl(3)
        for img in m['imagenes']:
            pdf.image(io.BytesIO(_jpeg(img['png'])), w=ancho * 0.92, x=Align.C)
            nl(3)
        for t in m['tablas']:
            if not t['filas']:
                continue
            texto(t['nombre'], 10.5, 'B', (30, 30, 40))
            n = len(t['columnas'])
            estado_tabla()
            with pdf.table(col_widths=anchos(t['columnas'], t['filas']),
                           text_align=tuple(['LEFT'] + ['CENTER'] * (n - 1)),
                           line_height=5, padding=1,
                           headings_style=FontFace(emphasis='BOLD', color=(255, 255, 255), fill_color=marca)) as tabla:
                cab = tabla.row()
                for c in t['columnas']:
                    cab.cell(_t(c))
                for r in t['filas'][:PDF_MAX_FILAS]:
                    fila = tabla.row()
                    for v in r:
                        fila.cell('' if v is None else _t(v))
            if len(t['filas']) > PDF_MAX_FILAS:
                texto(f"... {len(t['filas']) - PDF_MAX_FILAS} filas más (detalle completo en Excel/CSV).", 8, 'I', (130, 130, 140), 5)
            nl(3)
        for nota in m['notas']:
            texto(nota, 8.5, 'I', (110, 110, 120), 5)

    return bytes(pdf.output()), MIME['pdf'], 'pdf'


GENERADORES = {'pdf': generar_pdf, 'xlsx': generar_xlsx, 'csv': generar_csv, 'png': generar_png}
